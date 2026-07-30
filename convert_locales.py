import json
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font


def flatten_json(data, module="", prefix=""):
    rows = []

    if not isinstance(data, dict):
        return rows

    for key, value in data.items():
        if module == "":
            rows.extend(flatten_json(value, key, ""))
        else:
            if isinstance(value, dict):
                new_prefix = f"{prefix}.{key}" if prefix else key
                rows.extend(flatten_json(value, module, new_prefix))
            else:
                translation_key = f"{prefix}.{key}" if prefix else key
                rows.append((module, translation_key, value))

    return rows


def auto_adjust_columns(ws):
    for column_cells in ws.columns:
        length = max(len(str(cell.value or "")) for cell in column_cells)
        letter = column_cells[0].column_letter
        ws.column_dimensions[letter].width = min(length + 4, 80)


def workbook_matches(output_file, language, rows):
    """
    Returns True if the existing workbook already contains exactly
    the same data as what we're about to generate.
    """
    if not output_file.exists():
        return False

    try:
        wb = load_workbook(output_file, data_only=True)
        ws = wb.active

        expected = [("Module", "Translation Key", language)]
        expected.extend(rows)

        actual = []

        for row in ws.iter_rows(values_only=True):
            actual.append(tuple("" if v is None else v for v in row))

        expected = [
            tuple("" if v is None else v for v in row)
            for row in expected
        ]

        return actual == expected

    except Exception:
        return False


def convert_file(json_file, output_dir):
    language = json_file.stem.upper()

    print(f"Processing {json_file.name}...")

    with open(json_file, "r", encoding="utf-8") as f:
        data = json.load(f)

    rows = flatten_json(data)

    output_file = output_dir / f"{json_file.stem}.xlsx"

    # Skip rewriting if nothing changed
    if workbook_matches(output_file, language, rows):
        print(f"✓ {output_file.name} unchanged")
        return

    wb = Workbook()
    ws = wb.active
    ws.title = language

    headers = ["Module", "Translation Key", language]
    ws.append(headers)

    for cell in ws[1]:
        cell.font = Font(bold=True)

    for row in rows:
        ws.append(row)

    ws.freeze_panes = "A2"

    auto_adjust_columns(ws)

    wb.save(output_file)

    print(f"✓ Created {output_file.name}")


def main():
    if len(sys.argv) != 2:
        print("Usage:")
        print("  python convert_locales.py <project_name>")
        print()
        print("Example:")
        print("  python convert_locales.py setbet")
        sys.exit(1)

    project = sys.argv[1]

    base_dir = Path(__file__).parent

    project_dir = base_dir / project
    input_dir = project_dir / "locales-json"
    output_dir = project_dir / "locales-excel"

    if not project_dir.exists():
        print(f"❌ Project '{project}' not found.")
        sys.exit(1)

    if not input_dir.exists():
        print(f"❌ '{input_dir}' does not exist.")
        sys.exit(1)

    output_dir.mkdir(parents=True, exist_ok=True)

    json_files = sorted(input_dir.glob("*.json"))

    if not json_files:
        print("❌ No JSON files found.")
        sys.exit(1)

    print(f"\nProject : {project}")
    print(f"Input   : {input_dir}")
    print(f"Output  : {output_dir}")
    print(f"Languages Found : {len(json_files)}\n")

    for json_file in json_files:
        convert_file(json_file, output_dir)

    print("\n✅ Conversion completed successfully.")


if __name__ == "__main__":
    main()